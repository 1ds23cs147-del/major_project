from __future__ import annotations

import csv
import hashlib
import shutil
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path("project_data/datasets/rgb_thermal")
OUTPUT = Path("project_data/artifacts/clean_candidate_datasets")
AUDIT = Path("project_data/artifacts/leakage_audit")
IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".bmp", ".tif", ".tiff", ".webp"}


def paths_from_csv(path: Path) -> set[Path]:
    with path.open(newline="", encoding="utf-8") as handle:
        return {Path(row["path"]).resolve() for row in csv.DictReader(handle)}


def paths_from_workbook(path: Path, column: str, sheet_names: tuple[str, ...] | None = None) -> set[Path]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    values = set()
    names = sheet_names or tuple(workbook.sheetnames)
    for name in names:
        sheet = workbook[name]
        headers = [cell.value for cell in next(sheet.iter_rows())]
        if column not in headers:
            continue
        index = headers.index(column)
        for row in sheet.iter_rows(min_row=2, values_only=True):
            value = row[index]
            if value and value != "none":
                values.add(Path(value).resolve())
    workbook.close()
    return values


def copy_file(source: Path, destination: Path) -> None:
    destination.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source, destination)


def image_hashes(folder: Path) -> set[str]:
    hashes = set()
    for image in folder.iterdir():
        if image.is_file() and image.suffix.lower() in IMAGE_EXTENSIONS:
            hashes.add(hashlib.sha256(image.read_bytes()).hexdigest())
    return hashes


def build_rgb() -> dict:
    source = ROOT / "fasdd_yolo_real"
    destination = OUTPUT / "rgb_clean_candidate"
    if destination.exists():
        shutil.rmtree(destination)
    exact = paths_from_csv(AUDIT / "rgb_exact_cross_split_validation_quarantine.csv")
    near = paths_from_workbook(AUDIT / "rgb_near_duplicate_review.xlsx", "image_2")
    test_hashes = image_hashes(source / "test" / "images")
    test_collisions = set()
    for split in ("train", "val"):
        for image in (source / split / "images").iterdir():
            if image.is_file() and image.suffix.lower() in IMAGE_EXTENSIONS:
                if hashlib.sha256(image.read_bytes()).hexdigest() in test_hashes:
                    test_collisions.add(image.resolve())
    excluded = exact | near | test_collisions
    counts = {}
    for split in ("train", "val", "test"):
        copied = 0
        for image in sorted((source / split / "images").iterdir()):
            if image.suffix.lower() not in IMAGE_EXTENSIONS:
                continue
            if split in ("train", "val") and image.resolve() in excluded:
                continue
            copy_file(image, destination / split / "images" / image.name)
            label = source / split / "labels" / f"{image.stem}.txt"
            if label.exists():
                copy_file(label, destination / split / "labels" / label.name)
            copied += 1
        counts[split] = copied
    (destination / "data.yaml").write_text(
        f"path: {destination.resolve().as_posix()}\ntrain: train/images\nval: val/images\ntest: test/images\nnc: 2\nnames: [fire, smoke]\n",
        encoding="utf-8",
    )
    return {"source": str(source), "output": str(destination), "excluded_validation": len(excluded), "counts": counts}


def build_thermal() -> dict:
    source = ROOT / "flame3_cls_real" / "thermal"
    destination = OUTPUT / "thermal_clean_candidate"
    if destination.exists():
        shutil.rmtree(destination)
    near = paths_from_workbook(
        AUDIT / "thermal_duplicate_and_leakage_report.xlsx",
        "image_2",
        ("Thermal_Train_Val_Near",),
    )
    counts = {}
    for split in ("train", "val"):
        copied = 0
        for class_dir in sorted((source / split).iterdir()):
            if not class_dir.is_dir():
                continue
            for image in sorted(class_dir.iterdir()):
                if image.suffix.lower() not in IMAGE_EXTENSIONS:
                    continue
                if split == "val" and image.resolve() in near:
                    continue
                copy_file(image, destination / split / class_dir.name / image.name)
                copied += 1
        counts[split] = copied
    (destination / "data.yaml").write_text(
        f"path: {destination.resolve().as_posix()}\ntrain: train\nval: val\nnc: 2\nnames: [fire, no_fire]\n",
        encoding="utf-8",
    )
    return {"source": str(source), "output": str(destination), "excluded_validation": len(near), "counts": counts}


OUTPUT.mkdir(parents=True, exist_ok=True)
summary = {"rgb": build_rgb(), "thermal": build_thermal()}
(OUTPUT / "clean_split_manifest.txt").write_text(repr(summary) + "\n", encoding="utf-8")
print(summary)
